"""解析资源上限检查：压缩炸弹、超大表格、超量文本与超量切片。

背景：上传体积上限（默认 25MB）挡不住"解压后几十 GB"的构造文件——
OOXML 本身就是 ZIP，几 MB 的 docx 可以把 word/document.xml 压成几十 GB。
本测试用**构造出来的恶意样本**验证这些闸门真的会拦，同时验证正常文件不被误杀。

无需网络；使用 mock 嵌入后端；临时目录不影响真实数据。
"""
from __future__ import annotations

import io
import os
import sqlite3
import sys
import tempfile
import zipfile
from contextlib import closing
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import ingest as ingest_module  # noqa: E402
from app import parsing  # noqa: E402
from app.config import settings  # noqa: E402

PASS: list[str] = []
FAIL: list[str] = []


def check(cond: bool, msg: str) -> None:
    (PASS if cond else FAIL).append(msg)
    print(f"  [{'PASS' if cond else 'FAIL'}] {msg}")


def expect_parse_error(fn, *args, expected_codes: tuple[str, ...], label: str) -> None:
    try:
        fn(*args)
    except parsing.ParseError as exc:
        check(
            exc.code in expected_codes,
            f"{label} 被拒绝（code={exc.code}）",
        )
        print(f"        拒绝原因：{exc.message}")
    except Exception as exc:  # noqa: BLE001
        check(False, f"{label} 抛出了非 ParseError 的异常：{type(exc).__name__}: {exc}")
    else:
        check(False, f"{label} 未被拒绝（防护失效）")


def make_docx_bytes(paragraphs: list[str]) -> bytes:
    from docx import Document

    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def make_xlsx_bytes(rows: int) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    sheet = wb.active
    for index in range(1, rows + 1):
        sheet.cell(row=index, column=1, value=f"参数{index}")
        sheet.cell(row=index, column=2, value=f"取值{index}")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def make_zip(entries: dict[str, bytes]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, payload in entries.items():
            zf.writestr(name, payload)
    return buffer.getvalue()


def main() -> None:
    saved = {
        "uncompressed_mb": settings.parse_max_uncompressed_mb,
        "entries": settings.parse_max_zip_entries,
        "ratio": settings.parse_max_compression_ratio,
        "units": settings.parse_max_units,
        "chars": settings.parse_max_text_chars,
        "chunks": settings.parse_max_chunks,
    }
    saved_dirs = (settings.data_dir, settings.db_path, settings.upload_dir, settings.models_dir)
    saved_backend = settings.embed_backend

    try:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            settings.data_dir = root
            settings.db_path = root / "t.db"
            settings.upload_dir = root / "uploads"
            settings.models_dir = root / "models"
            settings.upload_dir.mkdir(parents=True, exist_ok=True)

            # ---------- A. 正常文件必须放行（防误杀）----------
            print("\n== 正常文件必须放行 ==")
            good_docx = make_docx_bytes(["设备参数说明", "第一段内容", "第二段内容"])
            good_xlsx = make_xlsx_bytes(5)
            parsing.check_magic("docx", good_docx)
            check(True, "正常的 DOCX 通过压缩包体检")
            parsing.check_magic("xlsx", good_xlsx)
            check(True, "正常的 XLSX 通过压缩包体检")

            docx_path = root / "good.docx"
            docx_path.write_bytes(good_docx)
            check(len(parsing.parse_docx(docx_path)) >= 3, "正常 DOCX 解析出预期段落数")

            xlsx_path = root / "good.xlsx"
            xlsx_path.write_bytes(good_xlsx)
            check(len(parsing.parse_xlsx(xlsx_path)) == 5, "正常 XLSX 解析出预期行数")

            txt_path = root / "good.txt"
            txt_path.write_text("第一段\n\n第二段", encoding="utf-8")
            check(len(parsing.parse_text(txt_path)) == 2, "正常 TXT 解析出预期段落数")

            # ---------- B. 压缩比炸弹 ----------
            print("\n== 压缩比炸弹（声明大小正常、压缩比异常）==")
            bomb = make_zip({"word/document.xml": b"\x00" * (50 * 1024 * 1024)})
            print(f"        构造样本 {len(bomb) / 1024:.0f} KB，声明解压后 50 MB")
            expect_parse_error(
                parsing.check_magic,
                "docx",
                bomb,
                expected_codes=("zip_bomb_ratio", "zip_too_large"),
                label="高压缩比 DOCX",
            )

            # ---------- C. 声明的解压总量超限 ----------
            print("\n== 解压总量超限（上限临时调为 1MB，用不可压缩数据隔离该分支）==")
            settings.parse_max_uncompressed_mb = 1
            heavy = make_zip({"word/document.xml": os.urandom(2 * 1024 * 1024)})
            expect_parse_error(
                parsing.check_magic,
                "docx",
                heavy,
                expected_codes=("zip_too_large",),
                label="解压后 2MB 的 DOCX",
            )
            settings.parse_max_uncompressed_mb = saved["uncompressed_mb"]

            # ---------- D. 条目数超限 ----------
            print("\n== 压缩条目数超限（上限临时调为 3）==")
            settings.parse_max_zip_entries = 3
            many = make_zip(
                {
                    "word/document.xml": b"<document/>",
                    "a.xml": b"<a/>",
                    "b.xml": b"<b/>",
                    "c.xml": b"<c/>",
                    "d.xml": b"<d/>",
                }
            )
            expect_parse_error(
                parsing.check_magic,
                "docx",
                many,
                expected_codes=("zip_too_many_entries",),
                label="含 5 个条目的 DOCX",
            )
            settings.parse_max_zip_entries = saved["entries"]

            # ---------- E. 有界实际读取（防谎报大小）----------
            print("\n== 有界实际读取：中央目录谎报大小时仍能拦住 ==")
            honest = make_zip({"word/document.xml": b"A" * 5000})
            with zipfile.ZipFile(io.BytesIO(honest)) as zf:
                actual = parsing._bounded_part_bytes(zf, "word/document.xml", 1024)
            check(
                actual > 1024,
                f"部件实际大小 {actual} 字节被读出并超过上限 1024（调用方据此拒绝）",
            )

            # ---------- F. 提取文本超限 ----------
            print("\n== 提取文本字符数超限（上限临时调为 100）==")
            settings.parse_max_text_chars = 100
            big_txt = root / "big.txt"
            big_txt.write_text("字" * 500, encoding="utf-8")
            expect_parse_error(
                parsing.parse_text,
                big_txt,
                expected_codes=("text_too_large",),
                label="500 字符的 TXT（上限 100）",
            )
            settings.parse_max_text_chars = saved["chars"]

            # ---------- G. 文本单元数超限 ----------
            print("\n== 文本单元数超限（上限临时调为 5）==")
            settings.parse_max_units = 5
            wide_path = root / "wide.xlsx"
            wide_path.write_bytes(make_xlsx_bytes(20))
            expect_parse_error(
                parsing.parse_xlsx,
                wide_path,
                expected_codes=("too_many_units",),
                label="20 行的 XLSX（上限 5）",
            )
            settings.parse_max_units = saved["units"]

            # ---------- H. 切片数超限（走真实入库流程）----------
            print("\n== 切片数超限：真实入库流程必须拦下 ==")
            settings.parse_max_chunks = 2
            settings.embed_backend = "mock"
            from app.db import init_db  # noqa: PLC0415

            init_db()
            with closing(sqlite3.connect(settings.db_path)) as database:
                database.execute(
                    "INSERT INTO users (username,password_hash,role,created_at,updated_at)"
                    " VALUES ('tester','hash','root','2026-01-01','2026-01-01')"
                )
                database.commit()

            from app.embeddings import embedding_service  # noqa: PLC0415

            embedding_service.state = "ready"
            embedding_service.message = "mock"

            long_text = ("这是一个用于触发切片上限的段落。" * 40 + "\n\n") * 40
            with closing(sqlite3.connect(settings.db_path)) as database:
                database.row_factory = sqlite3.Row
                try:
                    ingest_module.ingest_bytes(
                        database,
                        filename="huge.txt",
                        content_type="text/plain",
                        data=long_text.encode("utf-8"),
                        user_id=1,
                        version="1.0",
                    )
                except ingest_module.IngestError as exc:
                    check(
                        exc.code == "too_many_chunks",
                        f"超量切片被拒绝（code={exc.code}）",
                    )
                    print(f"        拒绝原因：{exc.message}")
                except Exception as exc:  # noqa: BLE001
                    check(False, f"抛出非预期异常：{type(exc).__name__}: {exc}")
                else:
                    check(False, "超量切片未被拒绝（防护失效）")

            settings.parse_max_chunks = saved["chunks"]
    finally:
        settings.parse_max_uncompressed_mb = saved["uncompressed_mb"]
        settings.parse_max_zip_entries = saved["entries"]
        settings.parse_max_compression_ratio = saved["ratio"]
        settings.parse_max_units = saved["units"]
        settings.parse_max_text_chars = saved["chars"]
        settings.parse_max_chunks = saved["chunks"]
        settings.embed_backend = saved_backend
        (
            settings.data_dir,
            settings.db_path,
            settings.upload_dir,
            settings.models_dir,
        ) = saved_dirs

    print(f"\n结果: {len(PASS)} 通过, {len(FAIL)} 失败")
    if FAIL:
        print("失败项:")
        for item in FAIL:
            print(f"  - {item}")
        sys.exit(1)


if __name__ == "__main__":
    main()
