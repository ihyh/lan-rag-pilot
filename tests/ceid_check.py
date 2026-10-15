"""CEID ingestion, exhaustive scoped retrieval and source-bound answer checks."""
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
_temp = tempfile.TemporaryDirectory(prefix="rag-ceid-check-")
os.environ.update(RAG_DATA_DIR=_temp.name, RAG_DB_PATH=str(Path(_temp.name)/"test.db"),
    RAG_UPLOAD_DIR=str(Path(_temp.name)/"uploads"), RAG_EMBED_BACKEND="mock",
    DEEPSEEK_API_KEY="isolated", DEEPSEEK_BASE_URL="http://127.0.0.1:19235/v1")
import numpy as np
from app import ceid, db as storage, llm
from app.chunking import chunk_units, TokenizerAdapter
from app.parsing import Unit
from app.index import VectorIndex
from app.routers.query import _select_sources

TABLE = """CEID Collection Event ID
11 = Load cycle started, tool platform 1
12 = Load cycle completed, tool platform 1
21 = Port Transfer State Ready to Load
21 = Load started, tool platform 2
22 = Load cycle completed, tool platform 2
40 = Pre-Load cycle started platform 1
111 = Unrelated event
Where used: S6F11
RCMD
11 = not an event
"""


class CeidChecks(unittest.TestCase):
    def setUp(self):
        self.db = sqlite3.connect(":memory:")
        self.db.row_factory = sqlite3.Row
        self.db.execute("PRAGMA foreign_keys=ON")
        self.db.executescript(storage.SCHEMA)
        self.index = VectorIndex()

    def tearDown(self):
        self.db.close()

    def add(self, did, text, filename="SECS.txt", similarity=.01):
        self.db.execute("INSERT INTO documents(id,filename,stored_name,content_type,size_bytes,sha256,status,created_at,updated_at) VALUES (?,?,?,?,?,?,'ready','now','now')",
                        (did,filename,str(did),"text/plain",len(text),str(did)))
        pieces = chunk_units([Unit(text, paragraph=1)], TokenizerAdapter(), 32, 4)
        for seq,p in enumerate(pieces):
            vector=np.array([similarity,np.sqrt(1-similarity**2)],dtype=np.float32)
            cur=self.db.execute("INSERT INTO chunks(document_id,seq,paragraph,token_count,content,vector) VALUES (?,?,?,?,?,?)",(did,seq,1,p.token_count,p.text,vector.tobytes()))
            if p.ceid_event:
                e=p.ceid_event
                self.db.execute("INSERT INTO ceid_events(document_id,chunk_id,ceid,event_name,scenario,full_definition,description_status,page_end) VALUES (?,?,?,?,?,?,?,?)",
                    (did,cur.lastrowid,e["ceid"],e["event_name"],e["scenario"],e["full_definition"],e["description_status"],e["page_end"]))
        self.index.reload(self.db)

    def sources(self, question, scope=None, k=2):
        hits=self.index.search(np.array([1,0],dtype=np.float32),k,scope,min_score=.9,query_text=question)
        sources=[]
        for h in hits:
            r=dict(self.db.execute("SELECT c.id AS chunk_id,c.document_id,c.seq,c.page,c.paragraph,c.content,d.filename FROM chunks c JOIN documents d ON d.id=c.document_id WHERE c.id=?",(h["chunk_id"],)).fetchone())
            sources.append(dict(r,score=h["score"],ceid_events=self.index.ceid_events_for(h["chunk_id"]),ceid_ids=self.index.ceid_ids_for(h["chunk_id"])))
        return _select_sources(sources,k,question)

    def test_case1_all_conflicts(self):
        self.add(1,TABLE)
        sources=self.sources("load cycle的事件ID为多少")
        answer=llm._source_projection("load cycle的事件ID为多少",sources)
        self.assertIn("11（tool platform 1",answer)
        self.assertIn("21（tool platform 2",answer)
        self.assertIn("Port Transfer State Ready to Load",answer)
        self.assertIn("port transfer",answer.lower())
        self.assertIn("同ID不同含义",answer)
        self.assertIn("人工甄别",answer)
        self.assertNotIn("更符合",answer)
        self.assertNotIn("CEID=12",answer)
        self.assertNotIn("CEID=40",answer)
        llm._validate_citations(answer,sources)

    def test_case2_incomplete_label_and_real_citation_index(self):
        self.add(1,TABLE)
        sources=self.sources("事件ID为11是什么")
        padded=[dict(chunk_id=-1,document_id=99,filename="irrelevant",content="x",ceid_events=[])]*2+sources
        answer=llm._source_projection("事件ID为11是什么",padded)
        self.assertIn("Load cycle started, tool platform 1",answer)
        self.assertIn("现有资料没有该条目的完整描述",answer)
        self.assertIn("[3]",answer)
        self.assertNotIn("CEID=111",answer)
        self.assertNotIn("CEID=21",answer)

    def test_atomic_long_cross_page_entry_and_metadata(self):
        first=Unit("CEID Collection Event ID\n11 = Load cycle started, tool platform 1\nDescription: "+"complete definition "*90,page=4)
        second=Unit("continuation remains in the same event.\n21 = Port Transfer State Ready to Load\nWhere used: S6F11",page=5)
        pieces=chunk_units([first,second],TokenizerAdapter(),32,4)
        entry=[p for p in pieces if p.ceid_event and p.ceid_event["ceid"]=="11"]
        self.assertEqual(len(entry),1)
        self.assertGreater(entry[0].token_count,32)
        self.assertIn("continuation remains in the same event.",entry[0].text)
        self.assertEqual(entry[0].ceid_event["page_end"],5)
        self.assertEqual(entry[0].ceid_event["description_status"],"provided")
        with self.assertRaises(Exception):chunk_units([first,second],TokenizerAdapter(),32,4,max_pieces=1)

    def test_all_same_id_documents_survive_k_and_score_acl(self):
        for did in range(1,19):self.add(did,f"CEID=21: Event in scenario {did}\nDescription: Definition {did}.")
        all_sources=self.sources("CEID=21是什么",k=1)
        self.assertEqual(len(all_sources),18)
        scoped=self.sources("CEID=21是什么",scope={1,2},k=1)
        self.assertEqual({s["document_id"] for s in scoped},{1,2})
        self.assertEqual(self.sources("CEID=21是什么",scope=set()),[])
        self.assertEqual(self.sources("CEID=210是什么"),[])

    def test_scenarios_and_other_numbers(self):
        self.assertEqual(ceid.event_spans("SVID\n11 = parameter\nRCMD=21"),[])
        self.assertEqual(ceid.event_spans("Power up note: CEID\n=\n5.\nMode CEID=3 if Auto\nCEID=80, n=slot count\nDescription: other message"),[])
        text="Below are the CEID description while the Access mode was disabled\n40 = Pre-Load cycle started platform 1\nBelow are the CEID description while the Access mode was enabled\n40 = E84 Transfer Blocked\nWhere used: S6F11"
        entries=ceid.event_spans(text)
        self.assertEqual(len(entries),2)
        self.assertIn("disabled",entries[0]["scenario"])
        self.assertIn("enabled",entries[1]["scenario"])
        self.assertNotIn("Below",entries[0]["full_definition"])

    def test_legacy_word_per_line_and_new_table_cascade(self):
        text="CEID\nCollection\nEvent\nID,\n1\nbyte;\nFormat:A5\n11\n=\nLoad\ncycle\nstarted,\ntool\nplatform\n1\n21\n=\nPort\nTransfer\nState\nReady\nto\nLoad\n21\n=\nLoad\nstarted,\ntool\nplatform\n2\nWhere\nused:\nS6F11"
        entries=ceid.event_spans(text)
        self.assertEqual([e["ceid"] for e in entries],["11","21","21"])
        self.assertEqual(entries[0]["event_name"],"Load cycle started, tool platform 1")
        self.add(1,text)
        self.db.execute("DELETE FROM chunks WHERE document_id=1")
        self.assertEqual(self.db.execute("SELECT COUNT(*) FROM ceid_events").fetchone()[0],0)

    def test_legacy_catalog_stops_before_other_field(self):
        self.db.executescript("DROP TABLE ceid_events;")
        self.db.execute("INSERT INTO documents(id,filename,stored_name,content_type,size_bytes,sha256,status,created_at,updated_at) VALUES (1,'old.pdf','old','pdf',1,'legacy','ready','now','now')")
        for seq,text in enumerate(("CEID Collection Event ID\n11 = Start\n21 = Ready\nWhere used: S6F11", "RCMD\n11 = command not event")):
            self.db.execute("INSERT INTO chunks(document_id,seq,token_count,content,vector) VALUES (1,?,?,?,?)",(seq,1,text,np.array([1,0],dtype=np.float32).tobytes()))
        self.index.reload(self.db)
        all_entries=[e for r in self.db.execute("SELECT id FROM chunks") for e in self.index.ceid_events_for(r[0])]
        self.assertEqual([e["event_name"] for e in all_entries],["Start","Ready"])

    def test_unknown_and_context_overflow_do_not_choose_subset(self):
        self.assertIn("没有可确认",ceid.answer("CEID=999是什么",[]))
        event=dict(ceid="21",event_name="Label",scenario="unknown",full_definition="Label",description_status="label_only")
        sources=[dict(chunk_id=i,content="X"*9000,ceid_events=[event]) for i in (1,2)]
        with self.assertRaises(ValueError):ceid.select_sources(sources,1,"CEID=21是什么")

    def test_exact_mentions_survive_top_k_without_becoming_definitions(self):
        self.add(1,TABLE)
        self.add(2,"Note: CEID=11 refers to a separate device observation.")
        sources=self.sources("事件ID为11是什么",k=1)
        self.assertEqual({s["document_id"] for s in sources},{1,2})
        answer=ceid.answer("事件ID为11是什么",sources)
        self.assertIn("其他同ID关联片段",answer)
        self.assertIn("separate device",answer)

    def test_explicit_docx_and_xlsx_table_rows(self):
        from docx import Document
        from openpyxl import Workbook
        from app.parsing import parse_docx,parse_xlsx
        headers=["CEID","事件名称","业务场景","完整定义"]
        values=["11","Load cycle started","tool platform 1","A full description. "*80]
        document=Document();table=document.add_table(rows=2,cols=4)
        for row,items in zip(table.rows,(headers,values)):
            for cell,value in zip(row.cells,items):cell.text=value
        docpath=Path(_temp.name)/"ceid.docx";document.save(docpath)
        workbook=Workbook();workbook.active.append(headers);workbook.active.append(values)
        xlsxpath=Path(_temp.name)/"ceid.xlsx";workbook.save(xlsxpath);workbook.close()
        for units in (parse_docx(docpath),parse_xlsx(xlsxpath)):
            pieces=chunk_units(units,TokenizerAdapter(),32,4)
            entries=[p for p in pieces if p.ceid_event]
            self.assertEqual(len(entries),1)
            self.assertEqual(entries[0].ceid_event["ceid"],"11")
            self.assertEqual(entries[0].ceid_event["scenario"],"tool platform 1")
            self.assertIn(values[3],entries[0].ceid_event["full_definition"])

    def test_chat_and_stream_do_not_delegate_conflict_arbitration(self):
        self.add(1,TABLE)
        sources=self.sources("load cycle的事件ID为多少")
        with patch.object(llm,"model_service_client",side_effect=AssertionError("CEID answer must preserve supplied catalog")):
            result=llm.chat("load cycle的事件ID为多少",sources)
            self.assertIn("Port Transfer",result["answer"])
            chunks=list(llm.stream_chat("load cycle的事件ID为多少",sources))
            self.assertTrue(chunks)


if __name__=="__main__":unittest.main()
